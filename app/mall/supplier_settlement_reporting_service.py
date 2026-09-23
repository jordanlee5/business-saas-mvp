"""供应商结算只读查询、已确认批次 Excel 导出与导出审计。"""

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_

from ..time_utils import UTC8_TIMEZONE, format_utc8, utc8_now
from .audit import MallAuditActionType
from .domain import SupplierSettlementStatus
from .supplier_settlement_service import (
    _time_key,
    _validate_confirmation_evidence,
    _validate_generation_evidence,
)


SETTLEMENT_STATUS_ALL = "ALL"
VALID_SETTLEMENT_STATUS_FILTERS = frozenset({
    SETTLEMENT_STATUS_ALL,
    SupplierSettlementStatus.PENDING_CONFIRMATION.value,
    SupplierSettlementStatus.CONFIRMED.value,
})
SETTLEMENT_STATUS_LABELS = {
    SupplierSettlementStatus.PENDING_CONFIRMATION.value: "待确认",
    SupplierSettlementStatus.CONFIRMED.value: "已确认",
}
SETTLEMENT_EXPORT_PERMISSION_MESSAGE = "当前账号无权导出供应商结算"
SETTLEMENT_EXPORT_STATE_MESSAGE = "仅已确认的供应商结算可以导出"
ZERO = Decimal("0.00")


@dataclass(frozen=True)
class SupplierSettlementListItem:
    batch_id: int
    settlement_public_id: str
    supplier_id: int
    supplier_public_id_snapshot: str
    supplier_name_snapshot: str
    status: str
    status_label: str
    period_start: datetime
    period_end: datetime
    order_count: int
    item_count: int
    total_quantity: int
    total_cost_amount: Decimal
    generated_by_admin_id: int
    generated_by_username: str
    generated_at: datetime
    confirmed_by_admin_id: int | None
    confirmed_by_username: str | None
    confirmed_at: datetime | None


@dataclass(frozen=True)
class SupplierSettlementPage:
    items: tuple[SupplierSettlementListItem, ...]
    keyword: str
    status: str
    page: int
    page_size: int
    total: int
    total_pages: int


@dataclass(frozen=True)
class SupplierSettlementDetailItem:
    settlement_item_id: int
    order_id: int
    order_item_id: int
    order_public_id_snapshot: str
    order_completed_at: datetime
    product_public_id_snapshot: str
    product_name_snapshot: str
    sku_code_snapshot: str
    sku_name_snapshot: str
    supplier_sku_code_snapshot: str | None
    unit_cost_price: Decimal
    quantity: int
    line_cost_amount: Decimal


@dataclass(frozen=True)
class SupplierSettlementDetail:
    batch_id: int
    settlement_public_id: str
    supplier_id: int
    supplier_public_id_snapshot: str
    supplier_name_snapshot: str
    status: str
    status_label: str
    period_start: datetime
    period_end: datetime
    order_count: int
    item_count: int
    total_quantity: int
    total_cost_amount: Decimal
    generated_by_admin_id: int
    generated_by_username: str
    generated_at: datetime
    confirmed_by_admin_id: int | None
    confirmed_by_username: str | None
    confirmed_at: datetime | None
    items: tuple[SupplierSettlementDetailItem, ...]


@dataclass(frozen=True)
class SupplierSettlementExportResult:
    batch_id: int
    settlement_public_id: str
    supplier_public_id_snapshot: str
    status: str
    total_cost_amount: Decimal
    exported_by_admin_id: int
    exported_at: datetime
    action_log_id: int
    filename: str
    content: bytes


def _integer(value, *, field_name, minimum, maximum=None):
    if (
        isinstance(value, bool)
        or not isinstance(value, int)
        or value < minimum
        or (maximum is not None and value > maximum)
    ):
        raise ValueError(f"{field_name}超出允许范围")
    return value


def _normalize_keyword(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError("查询关键词无效")
    normalized = " ".join(value.split())
    if len(normalized) > 100:
        raise ValueError("查询关键词不能超过100个字符")
    return normalized


def _normalize_status(value) -> str:
    if not isinstance(value, str):
        raise ValueError("供应商结算状态无效")
    normalized = value.strip().upper()
    if normalized not in VALID_SETTLEMENT_STATUS_FILTERS:
        raise ValueError("供应商结算状态无效")
    return normalized


def _normalize_public_id(value) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("供应商结算编号不能为空")
    normalized = value.strip()
    if len(normalized) > 32:
        raise ValueError("供应商结算编号不能超过 32 个字符")
    return normalized


def _wall_time(value, *, field_name):
    if not isinstance(value, datetime):
        raise RuntimeError(f"{field_name}无效")
    if value.tzinfo is not None and value.utcoffset() is not None:
        return value.astimezone(UTC8_TIMEZONE).replace(tzinfo=None)
    return value


def _optional_wall_time(value, *, field_name):
    if value is None:
        return None
    return _wall_time(value, field_name=field_name)


def _current_time(value):
    return _wall_time(
        utc8_now() if value is None else value,
        field_name="供应商结算导出时间",
    )


def _database_time(db, value):
    if db.get_bind().dialect.name == "postgresql":
        return value.replace(tzinfo=UTC8_TIMEZONE)
    return value


def _money(value, *, field_name):
    try:
        normalized = Decimal(value).quantize(Decimal("0.01"))
    except Exception as exc:
        raise RuntimeError(f"{field_name}无效") from exc
    if normalized < ZERO:
        raise RuntimeError(f"{field_name}无效")
    return normalized


def _status_label(value):
    label = SETTLEMENT_STATUS_LABELS.get(value)
    if label is None:
        raise RuntimeError("供应商结算状态异常")
    return label


def _admin_usernames(db, admin_ids):
    from ..models import User

    normalized_ids = tuple(sorted({
        admin_id for admin_id in admin_ids if admin_id is not None
    }))
    if not normalized_ids:
        return {}
    users = db.query(User).filter(User.id.in_(normalized_ids)).all()
    usernames = {user.id: user.username for user in users}
    if len(usernames) != len(normalized_ids):
        raise RuntimeError("供应商结算操作人证据不完整")
    return usernames


def _list_item(batch, usernames):
    generated_username = usernames.get(batch.generated_by_admin_id)
    if not generated_username:
        raise RuntimeError("供应商结算生成人证据不完整")
    confirmed_username = None
    if batch.confirmed_by_admin_id is not None:
        confirmed_username = usernames.get(batch.confirmed_by_admin_id)
        if not confirmed_username:
            raise RuntimeError("供应商结算确认人证据不完整")
    return SupplierSettlementListItem(
        batch_id=batch.id,
        settlement_public_id=batch.settlement_public_id,
        supplier_id=batch.supplier_id,
        supplier_public_id_snapshot=batch.supplier_public_id_snapshot,
        supplier_name_snapshot=batch.supplier_name_snapshot,
        status=batch.status,
        status_label=_status_label(batch.status),
        period_start=_wall_time(
            batch.period_start, field_name="结算开始时间"
        ),
        period_end=_wall_time(
            batch.period_end, field_name="结算结束时间"
        ),
        order_count=batch.order_count,
        item_count=batch.item_count,
        total_quantity=batch.total_quantity,
        total_cost_amount=_money(
            batch.total_cost_amount, field_name="结算成本总额"
        ),
        generated_by_admin_id=batch.generated_by_admin_id,
        generated_by_username=generated_username,
        generated_at=_wall_time(
            batch.generated_at, field_name="结算生成时间"
        ),
        confirmed_by_admin_id=batch.confirmed_by_admin_id,
        confirmed_by_username=confirmed_username,
        confirmed_at=_optional_wall_time(
            batch.confirmed_at, field_name="结算确认时间"
        ),
    )


def list_supplier_settlements(
    db,
    *,
    keyword="",
    status=SETTLEMENT_STATUS_ALL,
    page=1,
    page_size=20,
) -> SupplierSettlementPage:
    """分页查询结算批次快照；不修改状态或审计记录。"""
    from ..models import SupplierSettlementBatch

    _integer(page, field_name="页码", minimum=1)
    _integer(page_size, field_name="每页数量", minimum=1, maximum=100)
    normalized_keyword = _normalize_keyword(keyword)
    normalized_status = _normalize_status(status)

    with db.no_autoflush:
        query = db.query(SupplierSettlementBatch)
        if normalized_status != SETTLEMENT_STATUS_ALL:
            query = query.filter(
                SupplierSettlementBatch.status == normalized_status
            )
        if normalized_keyword:
            pattern = f"%{normalized_keyword}%"
            query = query.filter(or_(
                SupplierSettlementBatch.settlement_public_id.ilike(pattern),
                SupplierSettlementBatch.supplier_public_id_snapshot.ilike(
                    pattern
                ),
                SupplierSettlementBatch.supplier_name_snapshot.ilike(pattern),
            ))
        total = query.count()
        total_pages = max(1, (total + page_size - 1) // page_size)
        effective_page = min(page, total_pages)
        batches = tuple(
            query.order_by(
                SupplierSettlementBatch.generated_at.desc(),
                SupplierSettlementBatch.id.desc(),
            )
            .offset((effective_page - 1) * page_size)
            .limit(page_size)
            .all()
        )
        usernames = _admin_usernames(
            db,
            (
                admin_id
                for batch in batches
                for admin_id in (
                    batch.generated_by_admin_id,
                    batch.confirmed_by_admin_id,
                )
            ),
        )
        items = tuple(_list_item(batch, usernames) for batch in batches)

    return SupplierSettlementPage(
        items=items,
        keyword=normalized_keyword,
        status=normalized_status,
        page=effective_page,
        page_size=page_size,
        total=total,
        total_pages=total_pages,
    )


def get_supplier_settlement_detail(
    db,
    *,
    settlement_public_id,
) -> SupplierSettlementDetail:
    """读取并完整复核一个结算批次、来源订单和审计证据。"""
    from ..models import SupplierSettlementBatch

    normalized_public_id = _normalize_public_id(settlement_public_id)
    with db.no_autoflush:
        batch = db.query(SupplierSettlementBatch).filter(
            SupplierSettlementBatch.settlement_public_id
            == normalized_public_id
        ).one_or_none()
        if batch is None:
            raise ValueError("供应商结算批次不存在")
        items, _generation_log = _validate_generation_evidence(
            db, batch=batch
        )
        is_confirmed = (
            batch.status == SupplierSettlementStatus.CONFIRMED.value
        )
        _validate_confirmation_evidence(
            db,
            batch=batch,
            expect_confirmed=is_confirmed,
        )
        usernames = _admin_usernames(
            db,
            (
                batch.generated_by_admin_id,
                batch.confirmed_by_admin_id,
            ),
        )
        summary = _list_item(batch, usernames)
        detail_items = tuple(
            SupplierSettlementDetailItem(
                settlement_item_id=item.id,
                order_id=item.order_id,
                order_item_id=item.order_item_id,
                order_public_id_snapshot=item.order_public_id_snapshot,
                order_completed_at=_wall_time(
                    item.order_completed_at,
                    field_name="订单完成时间",
                ),
                product_public_id_snapshot=(
                    item.product_public_id_snapshot
                ),
                product_name_snapshot=item.product_name_snapshot,
                sku_code_snapshot=item.sku_code_snapshot,
                sku_name_snapshot=item.sku_name_snapshot,
                supplier_sku_code_snapshot=(
                    item.supplier_sku_code_snapshot
                ),
                unit_cost_price=_money(
                    item.unit_cost_price, field_name="单位成本"
                ),
                quantity=item.quantity,
                line_cost_amount=_money(
                    item.line_cost_amount, field_name="行成本"
                ),
            )
            for item in items
        )

    return SupplierSettlementDetail(
        **summary.__dict__,
        items=detail_items,
    )


def _excel_text(value) -> str:
    text = str(value or "")
    if text.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


def _append_table(sheet, headers, rows, *, money_columns=()):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="2563EB")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    money_indices = {
        index + 1 for index, name in enumerate(headers)
        if name in money_columns
    }
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(length + 2, 12), 42)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif cell.column in money_indices and isinstance(
                cell.value, (int, float)
            ):
                cell.number_format = "#,##0.00"


def build_supplier_settlement_workbook(
    detail: SupplierSettlementDetail,
    *,
    exported_at=None,
) -> BytesIO:
    """生成已确认结算的汇总及订单项成本快照工作簿。"""
    if detail.status != SupplierSettlementStatus.CONFIRMED.value:
        raise ValueError(SETTLEMENT_EXPORT_STATE_MESSAGE)
    export_time = _current_time(exported_at)
    workbook = Workbook()
    workbook.properties.creator = "business-saas-mvp"
    workbook.properties.title = (
        f"供应商结算 {detail.settlement_public_id}"
    )

    summary = workbook.active
    summary.title = "结算汇总"
    summary_rows = [
        ("结算编号", _excel_text(detail.settlement_public_id)),
        ("供应商公开编号", _excel_text(
            detail.supplier_public_id_snapshot
        )),
        ("供应商名称", _excel_text(detail.supplier_name_snapshot)),
        ("结算状态", detail.status_label),
        ("结算开始时间", detail.period_start),
        ("结算结束时间", detail.period_end),
        ("订单数", detail.order_count),
        ("订单项数", detail.item_count),
        ("商品数量", detail.total_quantity),
        ("人民币成本总额", float(detail.total_cost_amount)),
        ("生成人", _excel_text(detail.generated_by_username)),
        ("生成时间", detail.generated_at),
        ("确认人", _excel_text(detail.confirmed_by_username or "")),
        ("确认时间", detail.confirmed_at),
        ("导出时间", export_time),
    ]
    _append_table(
        summary,
        ["项目", "内容"],
        summary_rows,
    )
    summary.cell(row=11, column=2).number_format = "#,##0.00"

    item_sheet = workbook.create_sheet("结算明细")
    item_headers = [
        "结算明细ID", "订单公开编号", "订单完成时间",
        "商品公开编号", "商品名称", "SKU编码", "SKU名称",
        "供应商货号", "单位成本", "数量", "行成本",
    ]
    item_rows = [
        (
            item.settlement_item_id,
            _excel_text(item.order_public_id_snapshot),
            item.order_completed_at,
            _excel_text(item.product_public_id_snapshot),
            _excel_text(item.product_name_snapshot),
            _excel_text(item.sku_code_snapshot),
            _excel_text(item.sku_name_snapshot),
            _excel_text(item.supplier_sku_code_snapshot or ""),
            float(item.unit_cost_price),
            item.quantity,
            float(item.line_cost_amount),
        )
        for item in detail.items
    ]
    _append_table(
        item_sheet,
        item_headers,
        item_rows,
        money_columns=("单位成本", "行成本"),
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _require_export_actor(db, actor_admin_id):
    from ..admin_permissions import can_perform_mall_audit_action
    from ..models import User

    if (
        isinstance(actor_admin_id, bool)
        or not isinstance(actor_admin_id, int)
        or actor_admin_id <= 0
    ):
        raise PermissionError(SETTLEMENT_EXPORT_PERMISSION_MESSAGE)
    with db.no_autoflush:
        actor = db.query(User).filter(User.id == actor_admin_id).one_or_none()
    if (
        actor is None
        or actor.is_active is not True
        or not can_perform_mall_audit_action(
            actor,
            MallAuditActionType.SUPPLIER_SETTLEMENT_EXPORT,
        )
    ):
        raise PermissionError(SETTLEMENT_EXPORT_PERMISSION_MESSAGE)
    return actor


def record_supplier_settlement_export(
    db,
    *,
    actor_admin_id,
    detail: SupplierSettlementDetail,
    now=None,
):
    """记录一次已确认结算导出；调用方负责提交或回滚。"""
    from ..models import AdminActionLog

    if detail.status != SupplierSettlementStatus.CONFIRMED.value:
        raise ValueError(SETTLEMENT_EXPORT_STATE_MESSAGE)
    actor = _require_export_actor(db, actor_admin_id)
    export_time = _current_time(now)
    action_log = AdminActionLog(
        admin_id=actor.id,
        action_type=MallAuditActionType.SUPPLIER_SETTLEMENT_EXPORT.value,
        target_type="supplier_settlement_batch",
        target_id=detail.batch_id,
        description=(
            f"导出供应商结算 {detail.settlement_public_id}；"
            f"供应商 {detail.supplier_public_id_snapshot}；"
            f"订单 {detail.order_count}；明细 {detail.item_count}；"
            f"数量 {detail.total_quantity}；"
            f"成本 {detail.total_cost_amount:.2f}；"
            f"导出时间 {format_utc8(export_time)}"
        ),
        created_at=_database_time(db, export_time),
    )
    db.add(action_log)
    db.flush()
    return action_log


def execute_supplier_settlement_export(
    engine,
    *,
    actor_admin_id,
    settlement_public_id,
    now=None,
) -> SupplierSettlementExportResult:
    """在独立事务中复核、生成 Excel 并记录导出审计。"""
    from sqlalchemy.orm import Session

    export_time = _current_time(now)
    with engine.connect() as connection:
        if engine.dialect.name == "sqlite":
            connection.exec_driver_sql("BEGIN IMMEDIATE")
        else:
            connection.begin()
        try:
            with Session(
                bind=connection,
                autoflush=False,
                expire_on_commit=False,
            ) as db:
                _require_export_actor(db, actor_admin_id)
                detail = get_supplier_settlement_detail(
                    db,
                    settlement_public_id=settlement_public_id,
                )
                workbook = build_supplier_settlement_workbook(
                    detail,
                    exported_at=export_time,
                )
                action_log = record_supplier_settlement_export(
                    db,
                    actor_admin_id=actor_admin_id,
                    detail=detail,
                    now=export_time,
                )
                result = SupplierSettlementExportResult(
                    batch_id=detail.batch_id,
                    settlement_public_id=detail.settlement_public_id,
                    supplier_public_id_snapshot=(
                        detail.supplier_public_id_snapshot
                    ),
                    status=detail.status,
                    total_cost_amount=detail.total_cost_amount,
                    exported_by_admin_id=actor_admin_id,
                    exported_at=export_time,
                    action_log_id=action_log.id,
                    filename=(
                        "supplier_settlement_"
                        f"{detail.settlement_public_id}_"
                        f"{export_time:%Y%m%d%H%M%S}.xlsx"
                    ),
                    content=workbook.getvalue(),
                )
                connection.commit()
                return result
        except Exception:
            connection.rollback()
            raise
