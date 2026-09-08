"""管理员侧会员积分只读查询、对账导出与导出审计。"""

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_

from ..time_utils import UTC8_TIMEZONE, format_utc8, utc8_now
from .audit import MallAuditActionType
from .domain import (
    PointsGrantStatus,
    PointsLedgerEntryType,
    _to_utc8_wall_time,
    normalize_points,
)
from .points_ledger_service import audit_points_account_balance


MEMBER_STATUS_ALL = "ALL"
MEMBER_STATUS_ACTIVE = "ACTIVE"
MEMBER_STATUS_INACTIVE = "INACTIVE"
VALID_MEMBER_STATUSES = frozenset(
    {MEMBER_STATUS_ALL, MEMBER_STATUS_ACTIVE, MEMBER_STATUS_INACTIVE}
)
ZERO = Decimal("0.00")
EXPORT_PERMISSION_MESSAGE = "当前账号无权导出会员积分"


GRANT_STATUS_LABELS = {
    PointsGrantStatus.ACTIVE.value: "有效",
    PointsGrantStatus.EXHAUSTED.value: "已用尽",
    PointsGrantStatus.EXPIRED.value: "已过期",
    PointsGrantStatus.FROZEN.value: "已冻结",
}


LEDGER_TYPE_LABELS = {
    PointsLedgerEntryType.GRANT.value: "积分入账",
    PointsLedgerEntryType.RESERVE.value: "积分预占",
    PointsLedgerEntryType.RELEASE.value: "释放预占",
    PointsLedgerEntryType.CONSUME.value: "积分消费",
    PointsLedgerEntryType.REFUND.value: "积分退回",
    PointsLedgerEntryType.EXPIRE.value: "积分过期",
    PointsLedgerEntryType.ADJUST.value: "人工调整",
}


@dataclass(frozen=True)
class MemberPointsListItem:
    member_id: int
    member_public_id: str
    is_active: bool
    created_at: datetime
    account_id: int | None
    available_points: Decimal
    reserved_points: Decimal
    account_version: int
    grant_count: int
    ledger_count: int
    nearest_expires_at: datetime | None
    is_consistent: bool

    @property
    def total_points(self) -> Decimal:
        return normalize_points(
            self.available_points + self.reserved_points
        )


@dataclass(frozen=True)
class MemberPointsPage:
    items: tuple[MemberPointsListItem, ...]
    keyword: str
    member_status: str
    page: int
    page_size: int
    total: int
    total_pages: int


@dataclass(frozen=True)
class MemberPointsGrantItem:
    grant_id: int
    business_no: str
    uploader_name: str
    customer_name: str
    customer_phone: str
    plate_number: str
    customer_name_masked: str
    customer_phone_masked: str
    plate_number_masked: str
    granted_points: Decimal
    available_points: Decimal
    reserved_points: Decimal
    activated_at: datetime
    expires_at: datetime
    stored_status: str
    status_label: str


@dataclass(frozen=True)
class MemberPointsLedgerItem:
    ledger_id: int
    grant_id: int
    business_no: str
    entry_type: str
    entry_type_label: str
    available_points_delta: Decimal
    reserved_points_delta: Decimal
    actor_username: str | None
    reason: str | None
    reference_type: str | None
    reference_id: str | None
    created_at: datetime


@dataclass(frozen=True)
class MemberPointsDetail:
    member_id: int
    member_public_id: str
    is_active: bool
    member_created_at: datetime
    account_id: int | None
    available_points: Decimal
    reserved_points: Decimal
    account_version: int
    ledger_available_points: Decimal
    ledger_reserved_points: Decimal
    is_consistent: bool
    inconsistent_grant_ids: tuple[int, ...]
    grants: tuple[MemberPointsGrantItem, ...]
    ledgers: tuple[MemberPointsLedgerItem, ...]
    as_of: datetime

    @property
    def total_points(self) -> Decimal:
        return normalize_points(
            self.available_points + self.reserved_points
        )


def _integer(value, name, minimum, maximum=None):
    if (
        isinstance(value, bool)
        or not isinstance(value, int)
        or value < minimum
        or (maximum is not None and value > maximum)
    ):
        raise ValueError(f"{name}超出允许范围")


def _current_time(now):
    return _to_utc8_wall_time(
        utc8_now() if now is None else now,
        field_name="会员积分查询时间",
    )


def _database_time(db, value):
    if db.get_bind().dialect.name == "postgresql":
        return value.replace(tzinfo=UTC8_TIMEZONE)
    return value


def _normalize_keyword(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError("查询关键词无效")
    keyword = " ".join(value.split())
    if len(keyword) > 100:
        raise ValueError("查询关键词不能超过100个字符")
    return keyword


def _normalize_member_status(value) -> str:
    if not isinstance(value, str):
        raise ValueError("会员状态无效")
    normalized = value.strip().upper()
    if normalized not in VALID_MEMBER_STATUSES:
        raise ValueError("会员状态无效")
    return normalized


def _mask_name(value) -> str:
    text = str(value or "").strip()
    if not text:
        return "-"
    if len(text) == 1:
        return "*"
    return text[0] + "*" * min(len(text) - 1, 3)


def _mask_phone(value) -> str:
    text = str(value or "").strip()
    if not text:
        return "-"
    if len(text) <= 4:
        return "*" * len(text)
    if len(text) <= 7:
        return text[:2] + "*" * (len(text) - 4) + text[-2:]
    return text[:3] + "*" * (len(text) - 7) + text[-4:]


def _mask_plate(value) -> str:
    text = str(value or "").strip()
    if not text:
        return "-"
    if len(text) <= 2:
        return text[0] + "*" if text else "-"
    return text[:2] + "*" * (len(text) - 3) + text[-1]


def _grant_status_label(grant, current_time) -> str:
    if (
        grant.status == PointsGrantStatus.ACTIVE.value
        and _to_utc8_wall_time(
            grant.expires_at,
            field_name="积分到期时间",
        ) <= current_time
    ):
        return "待到期处理"
    return GRANT_STATUS_LABELS.get(grant.status, "异常状态")


def list_member_points(
    db,
    *,
    keyword="",
    member_status=MEMBER_STATUS_ALL,
    page=1,
    page_size=10,
    now=None,
) -> MemberPointsPage:
    """按会员分页查询积分汇总；不处理到期、不修复余额。"""
    from ..models import (
        BusinessRecord,
        Member,
        PointsAccount,
        PointsGrant,
    )

    _integer(page, "页码", 1)
    _integer(page_size, "每页数量", 1, 50)
    normalized_keyword = _normalize_keyword(keyword)
    normalized_status = _normalize_member_status(member_status)
    current_time = _current_time(now)

    with db.no_autoflush:
        query = (
            db.query(Member.id)
            .outerjoin(
                PointsAccount,
                PointsAccount.member_id == Member.id,
            )
            .outerjoin(
                PointsGrant,
                PointsGrant.account_id == PointsAccount.id,
            )
            .outerjoin(
                BusinessRecord,
                BusinessRecord.id == PointsGrant.business_record_id,
            )
        )
        if normalized_status == MEMBER_STATUS_ACTIVE:
            query = query.filter(Member.is_active.is_(True))
        elif normalized_status == MEMBER_STATUS_INACTIVE:
            query = query.filter(Member.is_active.is_(False))
        if normalized_keyword:
            pattern = f"%{normalized_keyword}%"
            query = query.filter(or_(
                Member.member_public_id.ilike(pattern),
                BusinessRecord.public_business_no.ilike(pattern),
                BusinessRecord.business_no.ilike(pattern),
                BusinessRecord.name.ilike(pattern),
                BusinessRecord.phone.ilike(pattern),
                BusinessRecord.plate_number.ilike(pattern),
            ))
        query = query.distinct()
        total = query.count()
        total_pages = max(1, (total + page_size - 1) // page_size)
        effective_page = min(page, total_pages)
        member_ids = tuple(
            row.id
            for row in (
                query.order_by(Member.id.desc())
                .offset((effective_page - 1) * page_size)
                .limit(page_size)
                .all()
            )
        )
        members = {
            member.id: member
            for member in db.query(Member).filter(
                Member.id.in_(member_ids)
            ).all()
        } if member_ids else {}
        accounts = {
            account.member_id: account
            for account in db.query(PointsAccount).filter(
                PointsAccount.member_id.in_(member_ids)
            ).all()
        } if member_ids else {}
        items = []
        stored_now = _database_time(db, current_time)
        for member_id in member_ids:
            member = members[member_id]
            account = accounts.get(member_id)
            if account is None:
                items.append(MemberPointsListItem(
                    member_id=member.id,
                    member_public_id=member.member_public_id,
                    is_active=bool(member.is_active),
                    created_at=_to_utc8_wall_time(
                        member.created_at, "会员创建时间"
                    ),
                    account_id=None,
                    available_points=ZERO,
                    reserved_points=ZERO,
                    account_version=0,
                    grant_count=0,
                    ledger_count=0,
                    nearest_expires_at=None,
                    is_consistent=False,
                ))
                continue
            grants = db.query(PointsGrant).filter(
                PointsGrant.account_id == account.id
            ).all()
            audit = audit_points_account_balance(
                db, account_id=account.id
            )
            grant_ids = tuple(grant.id for grant in grants)
            from ..models import PointsLedgerEntry
            ledger_count = (
                db.query(PointsLedgerEntry.id)
                .filter(PointsLedgerEntry.grant_id.in_(grant_ids))
                .count()
                if grant_ids else 0
            )
            nearest = (
                db.query(PointsGrant.expires_at)
                .filter(
                    PointsGrant.account_id == account.id,
                    PointsGrant.status == PointsGrantStatus.ACTIVE.value,
                    PointsGrant.expires_at > stored_now,
                    or_(
                        PointsGrant.available_points > 0,
                        PointsGrant.reserved_points > 0,
                    ),
                )
                .order_by(PointsGrant.expires_at.asc())
                .first()
            )
            items.append(MemberPointsListItem(
                member_id=member.id,
                member_public_id=member.member_public_id,
                is_active=bool(member.is_active),
                created_at=_to_utc8_wall_time(
                    member.created_at, "会员创建时间"
                ),
                account_id=account.id,
                available_points=normalize_points(
                    account.available_points
                ),
                reserved_points=normalize_points(account.reserved_points),
                account_version=account.version or 0,
                grant_count=len(grants),
                ledger_count=ledger_count,
                nearest_expires_at=(
                    _to_utc8_wall_time(nearest.expires_at, "积分到期时间")
                    if nearest else None
                ),
                is_consistent=audit.is_consistent,
            ))

    return MemberPointsPage(
        items=tuple(items),
        keyword=normalized_keyword,
        member_status=normalized_status,
        page=effective_page,
        page_size=page_size,
        total=total,
        total_pages=total_pages,
    )


def get_member_points_detail(
    db, *, member_id: int, now=None
) -> MemberPointsDetail:
    """读取会员账户、独立积分批次和不可变流水。"""
    from ..models import (
        BusinessRecord,
        Member,
        PointsAccount,
        PointsGrant,
        PointsLedgerEntry,
        User,
    )

    _integer(member_id, "会员编号", 1)
    current_time = _current_time(now)
    with db.no_autoflush:
        member = db.query(Member).filter(
            Member.id == member_id
        ).one_or_none()
        if member is None:
            raise ValueError("会员不存在")
        account = db.query(PointsAccount).filter(
            PointsAccount.member_id == member.id
        ).one_or_none()
        if account is None:
            return MemberPointsDetail(
                member_id=member.id,
                member_public_id=member.member_public_id,
                is_active=bool(member.is_active),
                member_created_at=_to_utc8_wall_time(
                    member.created_at, "会员创建时间"
                ),
                account_id=None,
                available_points=ZERO,
                reserved_points=ZERO,
                account_version=0,
                ledger_available_points=ZERO,
                ledger_reserved_points=ZERO,
                is_consistent=False,
                inconsistent_grant_ids=(),
                grants=(),
                ledgers=(),
                as_of=current_time,
            )
        grant_rows = (
            db.query(PointsGrant, BusinessRecord, User.username)
            .join(
                BusinessRecord,
                BusinessRecord.id == PointsGrant.business_record_id,
            )
            .outerjoin(User, User.id == BusinessRecord.user_id)
            .filter(PointsGrant.account_id == account.id)
            .order_by(
                PointsGrant.expires_at.asc(),
                PointsGrant.id.asc(),
            )
            .all()
        )
        grants = tuple(
            MemberPointsGrantItem(
                grant_id=grant.id,
                business_no=business.display_business_no,
                uploader_name=uploader_name or "-",
                customer_name=str(business.name or ""),
                customer_phone=str(business.phone or ""),
                plate_number=str(business.plate_number or ""),
                customer_name_masked=_mask_name(business.name),
                customer_phone_masked=_mask_phone(business.phone),
                plate_number_masked=_mask_plate(business.plate_number),
                granted_points=normalize_points(grant.granted_points),
                available_points=normalize_points(grant.available_points),
                reserved_points=normalize_points(grant.reserved_points),
                activated_at=_to_utc8_wall_time(
                    grant.activated_at, "积分激活时间"
                ),
                expires_at=_to_utc8_wall_time(
                    grant.expires_at, "积分到期时间"
                ),
                stored_status=grant.status,
                status_label=_grant_status_label(grant, current_time),
            )
            for grant, business, uploader_name in grant_rows
        )
        business_by_grant = {
            grant.id: business.display_business_no
            for grant, business, _uploader_name in grant_rows
        }
        ledger_rows = (
            db.query(PointsLedgerEntry, User.username)
            .join(
                PointsGrant,
                PointsGrant.id == PointsLedgerEntry.grant_id,
            )
            .outerjoin(User, User.id == PointsLedgerEntry.actor_admin_id)
            .filter(PointsGrant.account_id == account.id)
            .order_by(
                PointsLedgerEntry.created_at.desc(),
                PointsLedgerEntry.id.desc(),
            )
            .all()
        )
        ledgers = tuple(
            MemberPointsLedgerItem(
                ledger_id=entry.id,
                grant_id=entry.grant_id,
                business_no=business_by_grant.get(entry.grant_id, "-"),
                entry_type=entry.entry_type,
                entry_type_label=LEDGER_TYPE_LABELS.get(
                    entry.entry_type, "未知类型"
                ),
                available_points_delta=normalize_points(
                    entry.available_points_delta
                ),
                reserved_points_delta=normalize_points(
                    entry.reserved_points_delta
                ),
                actor_username=actor_username,
                reason=entry.reason,
                reference_type=entry.reference_type,
                reference_id=entry.reference_id,
                created_at=_to_utc8_wall_time(
                    entry.created_at, "积分流水时间"
                ),
            )
            for entry, actor_username in ledger_rows
        )
        audit = audit_points_account_balance(
            db, account_id=account.id
        )

    return MemberPointsDetail(
        member_id=member.id,
        member_public_id=member.member_public_id,
        is_active=bool(member.is_active),
        member_created_at=_to_utc8_wall_time(
            member.created_at, "会员创建时间"
        ),
        account_id=account.id,
        available_points=normalize_points(account.available_points),
        reserved_points=normalize_points(account.reserved_points),
        account_version=account.version or 0,
        ledger_available_points=audit.ledger_balance.available_points,
        ledger_reserved_points=audit.ledger_balance.reserved_points,
        is_consistent=audit.is_consistent,
        inconsistent_grant_ids=audit.inconsistent_grant_ids,
        grants=grants,
        ledgers=ledgers,
        as_of=current_time,
    )


def _excel_text(value) -> str:
    text = str(value or "")
    if text.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


def _append_table(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fill = PatternFill("solid", fgColor="2563EB")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        length = max(
            len(str(cell.value or "")) for cell in column_cells
        )
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(length + 2, 12), 36)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(cell.value, (int, float)):
                if cell.column in {
                    index + 1 for index, name in enumerate(headers)
                    if "积分" in name
                }:
                    cell.number_format = "#,##0.00"


def build_member_points_workbook(
    detail: MemberPointsDetail, *, exported_at=None
) -> BytesIO:
    """生成单会员三工作表对账文件；不包含登录和安全凭据。"""
    export_time = _current_time(exported_at)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "会员汇总"
    summary_rows = [
        ("会员公开编号", _excel_text(detail.member_public_id)),
        ("会员状态", "启用" if detail.is_active else "停用"),
        ("会员创建时间", detail.member_created_at),
        ("积分账户编号", detail.account_id or ""),
        ("可用积分", float(detail.available_points)),
        ("预占积分", float(detail.reserved_points)),
        ("合计积分", float(detail.total_points)),
        ("账户版本", detail.account_version),
        ("流水重算可用积分", float(detail.ledger_available_points)),
        ("流水重算预占积分", float(detail.ledger_reserved_points)),
        ("账本一致性", "一致" if detail.is_consistent else "异常"),
        ("导出时间", export_time),
    ]
    _append_table(summary, ["项目", "内容"], summary_rows)
    for row_number in (6, 7, 8, 10, 11):
        summary.cell(row=row_number, column=2).number_format = "#,##0.00"

    grant_sheet = workbook.create_sheet("积分批次")
    grant_headers = [
        "批次ID", "来源业务单号", "上传方", "客户姓名", "手机号",
        "车牌号", "原始授予积分", "可用积分", "预占积分", "激活时间",
        "到期时间", "状态",
    ]
    grant_rows = [
        (
            item.grant_id,
            _excel_text(item.business_no),
            _excel_text(item.uploader_name),
            _excel_text(item.customer_name),
            _excel_text(item.customer_phone),
            _excel_text(item.plate_number),
            float(item.granted_points),
            float(item.available_points),
            float(item.reserved_points),
            item.activated_at,
            item.expires_at,
            item.status_label,
        )
        for item in detail.grants
    ]
    _append_table(grant_sheet, grant_headers, grant_rows)

    ledger_sheet = workbook.create_sheet("积分流水")
    ledger_headers = [
        "流水ID", "批次ID", "来源业务单号", "流水类型", "可用积分变化",
        "预占积分变化", "操作管理员", "原因", "来源类型", "来源编号", "发生时间",
    ]
    ledger_rows = [
        (
            item.ledger_id,
            item.grant_id,
            _excel_text(item.business_no),
            item.entry_type_label,
            float(item.available_points_delta),
            float(item.reserved_points_delta),
            _excel_text(item.actor_username or ""),
            _excel_text(item.reason or ""),
            _excel_text(item.reference_type or ""),
            _excel_text(item.reference_id or ""),
            item.created_at,
        )
        for item in detail.ledgers
    ]
    _append_table(ledger_sheet, ledger_headers, ledger_rows)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def record_member_points_export(
    db,
    *,
    actor_admin_id: int,
    detail: MemberPointsDetail,
    now=None,
):
    """记录已生成的单会员积分导出；调用方负责提交或回滚。"""
    from ..admin_permissions import can_perform_mall_audit_action
    from ..models import AdminActionLog, User

    if (
        isinstance(actor_admin_id, bool)
        or not isinstance(actor_admin_id, int)
        or actor_admin_id <= 0
    ):
        raise PermissionError(EXPORT_PERMISSION_MESSAGE)
    with db.no_autoflush:
        actor = db.query(User).filter(
            User.id == actor_admin_id
        ).one_or_none()
    if (
        actor is None
        or actor.is_active is not True
        or not can_perform_mall_audit_action(
            actor, MallAuditActionType.MEMBER_POINTS_EXPORT
        )
    ):
        raise PermissionError(EXPORT_PERMISSION_MESSAGE)
    export_time = _current_time(now)
    log = AdminActionLog(
        admin_id=actor.id,
        action_type=MallAuditActionType.MEMBER_POINTS_EXPORT.value,
        target_type="member",
        target_id=detail.member_id,
        description=(
            f"生成会员积分对账导出：会员 {detail.member_public_id}；"
            f"积分批次 {len(detail.grants)} 个；流水 {len(detail.ledgers)} 条；"
            f"账本一致性：{'一致' if detail.is_consistent else '异常'}；"
            f"生成时间：{format_utc8(export_time)}"
        ),
        created_at=_database_time(db, export_time),
    )
    db.add(log)
    db.flush()
    return log
