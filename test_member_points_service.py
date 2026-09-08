import tempfile
import unittest
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.admin_permissions import OPERATOR, PRIMARY_REVIEWER, SUPER_ADMIN
from app.database import Base
from app.mall import (
    EXPORT_PERMISSION_MESSAGE,
    MEMBER_STATUS_ACTIVE,
    MEMBER_STATUS_INACTIVE,
    adjust_points_grant,
    build_member_points_workbook,
    get_member_points_detail,
    list_member_points,
    record_initial_points_grant,
    record_member_points_export,
)
from app.models import (
    AdminActionLog,
    BusinessRecord,
    Member,
    MemberWechatBinding,
    PointsAccount,
    PointsGrant,
    PointsLedgerEntry,
    UploadBatch,
    User,
)


NOW = datetime(2026, 9, 8, 12, 0, 0)
ZERO = Decimal("0.00")


class MemberPointsServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "member-points.db"
        self.engine = create_engine(f"sqlite:///{self.path}")
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(
            bind=self.engine,
            autoflush=False,
            expire_on_commit=False,
        )
        self.db = self.Session()
        self.super_admin = self.add_user(
            "member-points-super", "admin", SUPER_ADMIN
        )
        self.operator = self.add_user(
            "member-points-operator", "admin", OPERATOR
        )
        self.reviewer = self.add_user(
            "member-points-reviewer", "admin", PRIMARY_REVIEWER
        )
        self.partner = self.add_user("member-points-partner", "partner")
        self.batch = UploadBatch(
            user_id=self.partner.id,
            filename="member-points.xlsx",
            total_rows=2,
            success_rows=2,
            failed_rows=0,
            acceptance_status="已承接",
            redemption_mode="MALL_REDEMPTION",
            claim_deadline=NOW + timedelta(days=30),
        )
        self.db.add(self.batch)
        self.db.flush()
        self.member = Member(
            member_public_id="MEMBER-POINTS-001",
            created_at=NOW - timedelta(days=60),
        )
        self.db.add(self.member)
        self.db.flush()
        self.db.add(MemberWechatBinding(
            member_id=self.member.id,
            wechat_app_id="wx-secret-app",
            openid="openid-never-export",
            unionid="unionid-never-export",
        ))
        self.account = PointsAccount(member_id=self.member.id)
        self.db.add(self.account)
        self.db.flush()
        self.later_grant = self.add_grant(
            public_no="BR-MEMBER-001",
            name="陈测试",
            phone="13812345678",
            plate="桂A12345",
            points="100.50",
            expires_at=NOW + timedelta(days=20),
        )
        self.earlier_grant = self.add_grant(
            public_no="BR-MEMBER-002",
            name="=公式客户",
            phone="13987654321",
            plate="桂B98765",
            points="49.50",
            expires_at=NOW + timedelta(days=10),
        )
        adjust_points_grant(
            self.db,
            grant_id=self.later_grant.id,
            delta_points="-0.50",
            actor_admin_id=self.super_admin.id,
            reason="=公式原因",
            idempotency_key="MEMBER-DETAIL-ADJUST",
            now=NOW,
        )
        self.inactive_member = Member(
            member_public_id="MEMBER-INACTIVE-002",
            is_active=False,
            created_at=NOW,
        )
        self.db.add(self.inactive_member)
        self.db.flush()
        self.db.add(PointsAccount(member_id=self.inactive_member.id))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.temp.cleanup()

    def add_user(self, username, role, admin_level=None, is_active=True):
        user = User(
            username=username,
            password_hash="test-only",
            role=role,
            admin_level=admin_level,
            is_active=is_active,
        )
        self.db.add(user)
        self.db.flush()
        return user

    def add_grant(self, *, public_no, name, phone, plate, points, expires_at):
        business = BusinessRecord(
            user_id=self.partner.id,
            batch_id=self.batch.id,
            business_no=public_no,
            public_business_no=public_no,
            name=name,
            phone=phone,
            plate_number=plate,
            points_amount=Decimal(points),
            bank_card="",
            redemption_mode="MALL_REDEMPTION",
            claim_status="ACTIVATED",
        )
        self.db.add(business)
        self.db.flush()
        grant = PointsGrant(
            account_id=self.account.id,
            business_record_id=business.id,
            granted_points=Decimal(points),
            available_points=ZERO,
            reserved_points=ZERO,
            activated_at=NOW - timedelta(days=20),
            expires_at=expires_at,
            status="ACTIVE",
        )
        self.db.add(grant)
        self.db.flush()
        record_initial_points_grant(
            self.db,
            grant=grant,
            idempotency_key=f"fixture-member-points:{grant.id}",
            reference_type="BUSINESS_RECORD",
            reference_id=business.public_business_no,
            now=grant.activated_at,
        )
        return grant

    def test_list_summarizes_accounts_and_supports_source_keyword(self):
        result = list_member_points(
            self.db,
            keyword="13812345678",
            now=NOW,
        )
        self.assertEqual(result.total, 1)
        item = result.items[0]
        self.assertEqual(item.member_public_id, "MEMBER-POINTS-001")
        self.assertEqual(item.available_points, Decimal("149.50"))
        self.assertEqual(item.reserved_points, ZERO)
        self.assertEqual(item.grant_count, 2)
        self.assertEqual(item.ledger_count, 3)
        self.assertEqual(item.nearest_expires_at, NOW + timedelta(days=10))
        self.assertTrue(item.is_consistent)

    def test_list_filters_status_clamps_page_and_validates_input(self):
        active = list_member_points(
            self.db,
            member_status=MEMBER_STATUS_ACTIVE,
            page=99,
        )
        self.assertEqual(active.total, 1)
        self.assertEqual(active.page, 1)
        inactive = list_member_points(
            self.db, member_status=MEMBER_STATUS_INACTIVE
        )
        self.assertEqual(inactive.total, 1)
        self.assertFalse(inactive.items[0].is_active)
        self.assertTrue(inactive.items[0].is_consistent)
        for kwargs, message in (
            ({"page_size": 51}, "每页数量"),
            ({"member_status": "UNKNOWN"}, "会员状态"),
            ({"keyword": "x" * 101}, "关键词"),
        ):
            with self.subTest(kwargs=kwargs):
                with self.assertRaisesRegex(ValueError, message):
                    list_member_points(self.db, **kwargs)

    def test_detail_is_read_only_ordered_masked_and_consistent(self):
        before = (
            self.db.query(PointsLedgerEntry).count(),
            self.account.available_points,
            self.account.version,
        )
        detail = get_member_points_detail(
            self.db, member_id=self.member.id, now=NOW
        )
        self.assertEqual(
            [item.grant_id for item in detail.grants],
            [self.earlier_grant.id, self.later_grant.id],
        )
        later = detail.grants[1]
        self.assertEqual(later.customer_name_masked, "陈**")
        self.assertEqual(later.customer_phone_masked, "138****5678")
        self.assertEqual(later.plate_number_masked, "桂A****5")
        self.assertEqual(later.customer_name, "陈测试")
        self.assertTrue(detail.is_consistent)
        self.assertEqual(detail.ledgers[0].entry_type, "ADJUST")
        self.assertEqual(
            before,
            (
                self.db.query(PointsLedgerEntry).count(),
                self.account.available_points,
                self.account.version,
            ),
        )

    def test_detail_marks_due_and_detects_cache_drift_without_repair(self):
        self.earlier_grant.expires_at = NOW
        self.account.available_points = Decimal("1.00")
        self.db.commit()
        detail = get_member_points_detail(
            self.db, member_id=self.member.id, now=NOW
        )
        self.assertEqual(detail.grants[0].status_label, "待到期处理")
        self.assertFalse(detail.is_consistent)
        self.assertEqual(detail.available_points, Decimal("1.00"))
        self.assertEqual(detail.ledger_available_points, Decimal("149.50"))
        self.assertEqual(self.account.available_points, Decimal("1.00"))

    def test_detail_handles_missing_member_and_missing_account(self):
        with self.assertRaisesRegex(ValueError, "会员不存在"):
            get_member_points_detail(self.db, member_id=999999)
        orphan = Member(member_public_id="MEMBER-NO-ACCOUNT")
        self.db.add(orphan)
        self.db.flush()
        detail = get_member_points_detail(self.db, member_id=orphan.id)
        self.assertIsNone(detail.account_id)
        self.assertFalse(detail.is_consistent)
        self.assertEqual(detail.grants, ())

    def test_workbook_has_exact_sheets_and_escapes_formula_text(self):
        detail = get_member_points_detail(
            self.db, member_id=self.member.id, now=NOW
        )
        output = build_member_points_workbook(
            detail, exported_at=NOW
        )
        workbook = load_workbook(BytesIO(output.getvalue()))
        self.assertEqual(
            workbook.sheetnames,
            ["会员汇总", "积分批次", "积分流水"],
        )
        grants = workbook["积分批次"]
        self.assertEqual(grants["D2"].value, "'=公式客户")
        self.assertEqual(grants.freeze_panes, "A2")
        ledgers = workbook["积分流水"]
        all_values = "\n".join(
            str(cell.value or "")
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        self.assertIn("'=公式原因", all_values)
        self.assertNotIn("openid-never-export", all_values)
        self.assertNotIn("unionid-never-export", all_values)
        self.assertNotIn("MEMBER-DETAIL-ADJUST", all_values)
        self.assertNotIn("幂等", all_values)
        self.assertEqual(ledgers.auto_filter.ref, ledgers.dimensions)

    def test_export_audit_allows_operations_roles_and_is_rollback_safe(self):
        detail = get_member_points_detail(
            self.db, member_id=self.member.id, now=NOW
        )
        initial_count = self.db.query(AdminActionLog).count()
        log = record_member_points_export(
            self.db,
            actor_admin_id=self.operator.id,
            detail=detail,
            now=NOW,
        )
        self.assertEqual(log.action_type, "mall_member_points_export")
        self.assertEqual(log.target_type, "member")
        self.assertEqual(log.target_id, self.member.id)
        self.assertIn("积分批次 2 个", log.description)
        self.db.rollback()
        self.assertEqual(self.db.query(AdminActionLog).count(), initial_count)
        log = record_member_points_export(
            self.db,
            actor_admin_id=self.super_admin.id,
            detail=detail,
            now=NOW,
        )
        self.assertIsNotNone(log.id)

    def test_export_audit_rejects_reviewers_disabled_and_missing_actors(self):
        detail = get_member_points_detail(
            self.db, member_id=self.member.id, now=NOW
        )
        disabled = self.add_user(
            "disabled-exporter", "admin", SUPER_ADMIN, False
        )
        for actor_id in (self.reviewer.id, disabled.id, 999999, True):
            with self.subTest(actor_id=actor_id):
                with self.assertRaisesRegex(
                    PermissionError, EXPORT_PERMISSION_MESSAGE
                ):
                    record_member_points_export(
                        self.db,
                        actor_admin_id=actor_id,
                        detail=detail,
                        now=NOW,
                    )


if __name__ == "__main__":
    unittest.main()
