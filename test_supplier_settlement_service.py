import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from threading import Barrier
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.admin_permissions import OPERATOR, PRIMARY_REVIEWER, SUPER_ADMIN
from app.database import Base
from app.mall import (
    SETTLEMENT_CONFIRM_PERMISSION_MESSAGE,
    SETTLEMENT_EXPORT_PERMISSION_MESSAGE,
    SETTLEMENT_EXPORT_STATE_MESSAGE,
    SETTLEMENT_GENERATE_PERMISSION_MESSAGE,
    build_supplier_settlement_workbook,
    execute_supplier_settlement_confirmation,
    execute_supplier_settlement_export,
    execute_supplier_settlement_generation,
    get_supplier_settlement_detail,
    list_supplier_settlements,
)
import app.mall.supplier_settlement_service as settlement_service
import app.mall.supplier_settlement_reporting_service as reporting_service
from app.models import (
    AdminActionLog,
    Member,
    Order,
    OrderItem,
    Product,
    ProductCategory,
    ProductSku,
    Supplier,
    SupplierSettlementBatch,
    SupplierSettlementItem,
    User,
)
from app.time_utils import UTC8_TIMEZONE


NOW = datetime(2026, 9, 21, 12, 0, 0)
PERIOD_START = datetime(2026, 9, 1, 0, 0, 0)
PERIOD_END = datetime(2026, 9, 8, 0, 0, 0)


class SupplierSettlementServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "supplier-settlement.db"
        self.engine = create_engine(f"sqlite:///{self.path}")
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(
            bind=self.engine,
            autoflush=False,
            expire_on_commit=False,
        )
        self.db = self.Session()

        self.operator = User(
            username="settlement-operator",
            password_hash="test-only",
            role="admin",
            admin_level=OPERATOR,
            is_active=True,
        )
        self.super_admin = User(
            username="settlement-super-admin",
            password_hash="test-only",
            role="admin",
            admin_level=SUPER_ADMIN,
            is_active=True,
        )
        self.reviewer = User(
            username="settlement-reviewer",
            password_hash="test-only",
            role="admin",
            admin_level=PRIMARY_REVIEWER,
            is_active=True,
        )
        self.inactive_operator = User(
            username="inactive-settlement-operator",
            password_hash="test-only",
            role="admin",
            admin_level=OPERATOR,
            is_active=False,
        )
        self.inactive_super_admin = User(
            username="inactive-settlement-super-admin",
            password_hash="test-only",
            role="admin",
            admin_level=SUPER_ADMIN,
            is_active=False,
        )
        self.member = Member(
            member_public_id="MEM-SETTLEMENT-001",
            is_active=True,
            created_at=NOW,
            updated_at=NOW,
        )
        self.category = ProductCategory(
            name="结算测试分类",
            slug="settlement-test",
            is_active=True,
            created_at=NOW,
            updated_at=NOW,
        )
        self.supplier = Supplier(
            supplier_public_id="SUP-SETTLEMENT-001",
            name="结算测试供应商",
            is_active=True,
            created_at=NOW,
            updated_at=NOW,
        )
        self.other_supplier = Supplier(
            supplier_public_id="SUP-SETTLEMENT-002",
            name="其他结算供应商",
            is_active=True,
            created_at=NOW,
            updated_at=NOW,
        )
        self.db.add_all([
            self.operator,
            self.super_admin,
            self.reviewer,
            self.inactive_operator,
            self.inactive_super_admin,
            self.member,
            self.category,
            self.supplier,
            self.other_supplier,
        ])
        self.db.flush()
        self.product = Product(
            product_public_id="PRD-SETTLEMENT-001",
            category_id=self.category.id,
            name="结算测试商品",
            status="PUBLISHED",
            published_at=NOW,
            created_at=NOW,
            updated_at=NOW,
        )
        self.other_product = Product(
            product_public_id="PRD-SETTLEMENT-002",
            category_id=self.category.id,
            name="其他供应商商品",
            status="PUBLISHED",
            published_at=NOW,
            created_at=NOW,
            updated_at=NOW,
        )
        self.db.add_all([self.product, self.other_product])
        self.db.flush()
        self.sku_one = self.add_sku(
            supplier=self.supplier,
            product=self.product,
            code="SETTLEMENT-SKU-001",
            name="标准款",
            cost="12.50",
        )
        self.sku_two = self.add_sku(
            supplier=self.supplier,
            product=self.product,
            code="SETTLEMENT-SKU-002",
            name="轻量款",
            cost="7.25",
        )
        self.other_sku = self.add_sku(
            supplier=self.other_supplier,
            product=self.other_product,
            code="SETTLEMENT-SKU-003",
            name="其他款",
            cost="9.00",
        )
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.temp.cleanup()

    def add_sku(self, *, supplier, product, code, name, cost):
        sku = ProductSku(
            product_id=product.id,
            supplier_id=supplier.id,
            sku_code=code,
            name=name,
            points_price=Decimal("50.00"),
            cost_price=Decimal(cost),
            low_stock_threshold=1,
            is_active=True,
            created_at=NOW,
            updated_at=NOW,
        )
        self.db.add(sku)
        self.db.flush()
        return sku

    def add_order(
        self,
        *,
        suffix,
        completed_at,
        lines,
        status="COMPLETED",
    ):
        total_points = sum(
            (Decimal("50.00") * quantity for _sku, quantity in lines),
            Decimal("0.00"),
        )
        total_cost = sum(
            (Decimal(sku.cost_price) * quantity for sku, quantity in lines),
            Decimal("0.00"),
        )
        total_quantity = sum(quantity for _sku, quantity in lines)
        shipped_at = None
        order_completed_at = None
        refund_reason = None
        refunded_at = None
        if status in ("SHIPPED", "COMPLETED", "REFUNDED"):
            base_time = completed_at or NOW
            shipped_at = base_time - timedelta(hours=1)
        if status in ("COMPLETED", "REFUNDED"):
            order_completed_at = completed_at
        if status == "REFUNDED":
            refund_reason = "结算排除退款订单"
            refunded_at = completed_at + timedelta(hours=1)

        order = Order(
            order_public_id=f"ORD-SETTLEMENT-{suffix}",
            idempotency_key=f"settlement-order-{suffix}",
            member_id=self.member.id,
            status=status,
            total_points=total_points,
            total_cost_amount=total_cost,
            total_quantity=total_quantity,
            shipping_carrier=("顺丰速运" if shipped_at else None),
            tracking_number=(f"SF-{suffix}" if shipped_at else None),
            shipped_at=shipped_at,
            completed_at=order_completed_at,
            refund_reason=refund_reason,
            refunded_at=refunded_at,
            created_at=NOW - timedelta(days=30),
            updated_at=NOW,
        )
        self.db.add(order)
        self.db.flush()
        items = []
        for sku, quantity in lines:
            supplier = (
                self.supplier
                if sku.supplier_id == self.supplier.id
                else self.other_supplier
            )
            product = (
                self.product
                if sku.product_id == self.product.id
                else self.other_product
            )
            unit_cost = Decimal(sku.cost_price)
            item = OrderItem(
                order_id=order.id,
                product_id=product.id,
                sku_id=sku.id,
                supplier_id=supplier.id,
                product_public_id_snapshot=product.product_public_id,
                product_name_snapshot=product.name,
                sku_code_snapshot=sku.sku_code,
                sku_name_snapshot=sku.name,
                supplier_public_id_snapshot=supplier.supplier_public_id,
                supplier_name_snapshot=supplier.name,
                supplier_sku_code_snapshot=f"SUP-{sku.sku_code}",
                product_image_path_snapshot=None,
                unit_points_price=Decimal("50.00"),
                unit_cost_price=unit_cost,
                quantity=quantity,
                line_points=Decimal("50.00") * quantity,
                line_cost_amount=unit_cost * quantity,
                created_at=NOW - timedelta(days=30),
            )
            self.db.add(item)
            items.append(item)
        self.db.flush()
        return order, tuple(items)

    def generate(self, **overrides):
        request = {
            "actor_admin_id": self.operator.id,
            "supplier_id": self.supplier.id,
            "period_start": PERIOD_START,
            "period_end": PERIOD_END,
            "now": NOW,
        }
        request.update(overrides)
        return execute_supplier_settlement_generation(
            self.engine,
            **request,
        )

    def confirm(self, settlement_public_id, **overrides):
        request = {
            "actor_admin_id": self.super_admin.id,
            "settlement_public_id": settlement_public_id,
            "now": NOW + timedelta(minutes=1),
        }
        request.update(overrides)
        return execute_supplier_settlement_confirmation(
            self.engine,
            **request,
        )

    def test_generates_only_unsettled_completed_items_in_half_open_period(self):
        start_order, start_items = self.add_order(
            suffix="START",
            completed_at=PERIOD_START,
            lines=((self.sku_one, 2), (self.sku_two, 1)),
        )
        middle_order, middle_items = self.add_order(
            suffix="MIDDLE",
            completed_at=PERIOD_START + timedelta(days=3),
            lines=((self.sku_one, 1), (self.other_sku, 1)),
        )
        self.add_order(
            suffix="END",
            completed_at=PERIOD_END,
            lines=((self.sku_one, 1),),
        )
        self.add_order(
            suffix="BEFORE",
            completed_at=PERIOD_START - timedelta(seconds=1),
            lines=((self.sku_one, 1),),
        )
        self.add_order(
            suffix="REFUNDED",
            completed_at=PERIOD_START + timedelta(days=2),
            lines=((self.sku_one, 1),),
            status="REFUNDED",
        )
        self.add_order(
            suffix="SHIPPED",
            completed_at=PERIOD_START + timedelta(days=2),
            lines=((self.sku_one, 1),),
            status="SHIPPED",
        )
        self.db.commit()

        original_supplier_name = self.supplier.name
        self.supplier.name = "结算时供应商名称"
        self.supplier.is_active = False
        self.db.commit()

        result = self.generate()

        self.assertEqual(result.status, "PENDING_CONFIRMATION")
        self.assertEqual(result.order_count, 2)
        self.assertEqual(result.item_count, 3)
        self.assertEqual(result.total_quantity, 4)
        self.assertEqual(result.total_cost_amount, Decimal("44.75"))
        self.assertEqual(result.supplier_name_snapshot, "结算时供应商名称")
        self.assertEqual(len(result.settlement_item_ids), 3)

        with self.Session() as db:
            stored_items = (
                db.query(SupplierSettlementItem)
                .order_by(SupplierSettlementItem.order_item_id.asc())
                .all()
            )
            self.assertEqual(
                {item.order_item_id for item in stored_items},
                {
                    start_items[0].id,
                    start_items[1].id,
                    middle_items[0].id,
                },
            )
            self.assertEqual(
                {item.order_id for item in stored_items},
                {start_order.id, middle_order.id},
            )
            self.assertEqual(
                {item.supplier_name_snapshot for item in stored_items},
                {original_supplier_name},
            )
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_generate",
                    target_type="supplier_settlement_batch",
                    target_id=result.batch_id,
                ).count(),
                1,
            )

    def test_generated_snapshots_do_not_follow_later_source_changes(self):
        order, items = self.add_order(
            suffix="IMMUTABLE",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        result = self.generate()

        self.db.expire_all()
        source_item = self.db.get(OrderItem, items[0].id)
        source_item.product_name_snapshot = "订单源快照后续被纠正"
        source_item.unit_cost_price = Decimal("20.00")
        source_item.line_cost_amount = Decimal("20.00")
        order = self.db.get(Order, order.id)
        order.order_public_id = "ORD-SOURCE-CHANGED"
        self.db.commit()

        with self.Session() as db:
            stored = db.query(SupplierSettlementItem).filter_by(
                settlement_batch_id=result.batch_id
            ).one()
            self.assertEqual(
                stored.order_public_id_snapshot,
                "ORD-SETTLEMENT-IMMUTABLE",
            )
            self.assertEqual(stored.product_name_snapshot, "结算测试商品")
            self.assertEqual(stored.unit_cost_price, Decimal("12.50"))
            self.assertEqual(stored.line_cost_amount, Decimal("12.50"))

    def test_repeated_generation_never_settles_an_order_item_twice(self):
        _order, first_items = self.add_order(
            suffix="FIRST",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        first = self.generate()

        with self.assertRaisesRegex(ValueError, "没有未结算"):
            self.generate()

        _order, second_items = self.add_order(
            suffix="SECOND",
            completed_at=PERIOD_START + timedelta(days=2),
            lines=((self.sku_two, 2),),
        )
        self.db.commit()
        second = self.generate()

        self.assertNotEqual(first.batch_id, second.batch_id)
        with self.Session() as db:
            stored = db.query(SupplierSettlementItem).all()
            self.assertEqual(len(stored), 2)
            self.assertEqual(
                {item.order_item_id for item in stored},
                {first_items[0].id, second_items[0].id},
            )

    def test_rejects_unauthorized_or_inactive_actor_without_writes(self):
        self.add_order(
            suffix="PERMISSION",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()

        for actor_id in (self.reviewer.id, self.inactive_operator.id, 0):
            with self.subTest(actor_id=actor_id):
                with self.assertRaisesRegex(
                    PermissionError,
                    SETTLEMENT_GENERATE_PERMISSION_MESSAGE,
                ):
                    self.generate(actor_admin_id=actor_id)
        with self.Session() as db:
            self.assertEqual(db.query(SupplierSettlementBatch).count(), 0)
            self.assertEqual(db.query(SupplierSettlementItem).count(), 0)

    def test_rejects_invalid_period_generation_time_and_supplier(self):
        self.add_order(
            suffix="VALIDATION",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()

        invalid_requests = (
            (
                {"period_start": PERIOD_END, "period_end": PERIOD_START},
                "必须晚于",
            ),
            ({"now": PERIOD_END - timedelta(seconds=1)}, "不能早于"),
            ({"supplier_id": 999999}, "供应商不存在"),
            ({"period_start": "2026-09-01"}, "开始时间无效"),
        )
        for overrides, message in invalid_requests:
            with self.subTest(overrides=overrides):
                with self.assertRaisesRegex(ValueError, message):
                    self.generate(**overrides)

    def test_accepts_aware_utc8_period_and_normalizes_for_sqlite(self):
        self.add_order(
            suffix="AWARE",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()

        result = self.generate(
            period_start=PERIOD_START.replace(tzinfo=UTC8_TIMEZONE),
            period_end=PERIOD_END.replace(tzinfo=UTC8_TIMEZONE),
            now=NOW.replace(tzinfo=UTC8_TIMEZONE),
        )

        self.assertIsNone(result.period_start.tzinfo)
        self.assertIsNone(result.period_end.tzinfo)
        self.assertEqual(result.period_start, PERIOD_START)
        self.assertEqual(result.period_end, PERIOD_END)

    def test_transaction_rolls_back_batch_items_and_log_on_late_failure(self):
        self.add_order(
            suffix="ROLLBACK",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()

        with patch.object(
            settlement_service,
            "_validate_generation_evidence",
            side_effect=RuntimeError("模拟结算证据核验失败"),
        ):
            with self.assertRaisesRegex(RuntimeError, "模拟结算"):
                self.generate()

        with self.Session() as db:
            self.assertEqual(db.query(SupplierSettlementBatch).count(), 0)
            self.assertEqual(db.query(SupplierSettlementItem).count(), 0)
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_generate"
                ).count(),
                0,
            )

    def test_confirms_batch_once_and_exact_replay_is_stable(self):
        self.add_order(
            suffix="CONFIRM",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 2),),
        )
        self.db.commit()
        generated = self.generate()

        first = self.confirm(generated.settlement_public_id)
        replay = self.confirm(generated.settlement_public_id)

        self.assertFalse(first.replayed)
        self.assertTrue(replay.replayed)
        self.assertEqual(first.status, "CONFIRMED")
        self.assertEqual(first.confirmed_by_admin_id, self.super_admin.id)
        self.assertEqual(first.confirmed_at, NOW + timedelta(minutes=1))
        self.assertEqual(replay.action_log_id, first.action_log_id)
        self.assertEqual(replay.confirmed_at, first.confirmed_at)
        with self.Session() as db:
            batch = db.get(SupplierSettlementBatch, generated.batch_id)
            self.assertEqual(batch.status, "CONFIRMED")
            self.assertEqual(batch.confirmed_by_admin_id, self.super_admin.id)
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_confirm",
                    target_type="supplier_settlement_batch",
                    target_id=batch.id,
                ).count(),
                1,
            )

    def test_concurrent_confirmation_creates_one_audit_evidence(self):
        self.add_order(
            suffix="CONCURRENT-CONFIRM",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()
        barrier = Barrier(2)

        def confirm_once():
            barrier.wait()
            return self.confirm(generated.settlement_public_id)

        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(confirm_once) for _ in range(2)]
            outcomes = [future.result() for future in futures]

        self.assertEqual(
            sorted(result.replayed for result in outcomes),
            [False, True],
        )
        with self.Session() as db:
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_confirm"
                ).count(),
                1,
            )

    def test_confirmation_requires_active_super_admin(self):
        self.add_order(
            suffix="CONFIRM-PERMISSION",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()

        for actor_id in (
            self.operator.id,
            self.reviewer.id,
            self.inactive_super_admin.id,
            0,
        ):
            with self.subTest(actor_id=actor_id):
                with self.assertRaisesRegex(
                    PermissionError,
                    SETTLEMENT_CONFIRM_PERMISSION_MESSAGE,
                ):
                    self.confirm(
                        generated.settlement_public_id,
                        actor_admin_id=actor_id,
                    )
        with self.Session() as db:
            batch = db.get(SupplierSettlementBatch, generated.batch_id)
            self.assertEqual(batch.status, "PENDING_CONFIRMATION")
            self.assertIsNone(batch.confirmed_at)

    def test_confirmation_rejects_invalid_time_missing_and_refunded_source(self):
        order, _items = self.add_order(
            suffix="CONFIRM-VALIDATION",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()

        with self.assertRaisesRegex(ValueError, "不能早于生成时间"):
            self.confirm(
                generated.settlement_public_id,
                now=NOW - timedelta(seconds=1),
            )
        with self.assertRaisesRegex(ValueError, "批次不存在"):
            self.confirm("STL-NOT-FOUND")

        self.db.expire_all()
        stored_order = self.db.get(Order, order.id)
        stored_order.status = "REFUNDED"
        stored_order.refund_reason = "结算确认前发生退款"
        stored_order.refunded_at = NOW + timedelta(seconds=30)
        self.db.commit()
        with self.assertRaisesRegex(RuntimeError, "退款证据异常"):
            self.confirm(generated.settlement_public_id)
        with self.Session() as db:
            batch = db.get(SupplierSettlementBatch, generated.batch_id)
            self.assertEqual(batch.status, "PENDING_CONFIRMATION")

    def test_confirmation_fails_closed_for_tampered_batch_evidence(self):
        self.add_order(
            suffix="CONFIRM-TAMPER",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()
        self.db.expire_all()
        batch = self.db.get(SupplierSettlementBatch, generated.batch_id)
        batch.total_cost_amount = Decimal("999.00")
        self.db.commit()

        with self.assertRaisesRegex(RuntimeError, "合计或状态证据不完整"):
            self.confirm(generated.settlement_public_id)
        with self.Session() as db:
            batch = db.get(SupplierSettlementBatch, generated.batch_id)
            self.assertEqual(batch.status, "PENDING_CONFIRMATION")
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_confirm"
                ).count(),
                0,
            )

    def test_confirmation_rolls_back_status_and_log_on_late_failure(self):
        self.add_order(
            suffix="CONFIRM-ROLLBACK",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()
        original_validator = settlement_service._validate_confirmation_evidence

        def fail_after_write(db, *, batch, expect_confirmed):
            if expect_confirmed:
                raise RuntimeError("模拟确认后证据核验失败")
            return original_validator(
                db,
                batch=batch,
                expect_confirmed=expect_confirmed,
            )

        with patch.object(
            settlement_service,
            "_validate_confirmation_evidence",
            side_effect=fail_after_write,
        ):
            with self.assertRaisesRegex(RuntimeError, "模拟确认后"):
                self.confirm(generated.settlement_public_id)

        with self.Session() as db:
            batch = db.get(SupplierSettlementBatch, generated.batch_id)
            self.assertEqual(batch.status, "PENDING_CONFIRMATION")
            self.assertIsNone(batch.confirmed_by_admin_id)
            self.assertIsNone(batch.confirmed_at)
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_confirm"
                ).count(),
                0,
            )

    def test_lists_settlements_with_keyword_status_and_pagination(self):
        self.add_order(
            suffix="REPORT-LIST-CONFIRMED",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        confirmed = self.generate()
        self.confirm(confirmed.settlement_public_id)

        self.add_order(
            suffix="REPORT-LIST-PENDING",
            completed_at=PERIOD_START + timedelta(days=2),
            lines=((self.sku_two, 2),),
        )
        self.db.commit()
        pending = self.generate(now=NOW + timedelta(minutes=2))

        with self.Session() as db:
            page = list_supplier_settlements(db, page=1, page_size=1)
            self.assertEqual(page.total, 2)
            self.assertEqual(page.total_pages, 2)
            self.assertEqual(page.items[0].batch_id, pending.batch_id)
            self.assertEqual(page.items[0].status_label, "待确认")
            self.assertEqual(
                page.items[0].generated_by_username,
                self.operator.username,
            )

            confirmed_page = list_supplier_settlements(
                db,
                status="CONFIRMED",
            )
            self.assertEqual(confirmed_page.total, 1)
            self.assertEqual(
                confirmed_page.items[0].settlement_public_id,
                confirmed.settlement_public_id,
            )
            self.assertEqual(
                confirmed_page.items[0].confirmed_by_username,
                self.super_admin.username,
            )

            keyword_page = list_supplier_settlements(
                db,
                keyword=confirmed.settlement_public_id,
            )
            self.assertEqual(keyword_page.total, 1)
            supplier_page = list_supplier_settlements(
                db,
                keyword=self.supplier.name,
            )
            self.assertEqual(supplier_page.total, 2)

            for request in (
                {"status": "UNKNOWN"},
                {"page": 0},
                {"page_size": 101},
                {"keyword": "x" * 101},
            ):
                with self.subTest(request=request):
                    with self.assertRaises(ValueError):
                        list_supplier_settlements(db, **request)

    def test_detail_revalidates_source_and_returns_snapshots(self):
        _order, order_items = self.add_order(
            suffix="REPORT-DETAIL",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 2), (self.sku_two, 1)),
        )
        self.db.commit()
        generated = self.generate()
        self.confirm(generated.settlement_public_id)

        with self.Session() as db:
            detail = get_supplier_settlement_detail(
                db,
                settlement_public_id=generated.settlement_public_id,
            )
            self.assertEqual(detail.status, "CONFIRMED")
            self.assertEqual(detail.status_label, "已确认")
            self.assertEqual(detail.order_count, 1)
            self.assertEqual(detail.item_count, 2)
            self.assertEqual(detail.total_quantity, 3)
            self.assertEqual(detail.total_cost_amount, Decimal("32.25"))
            self.assertEqual(len(detail.items), 2)
            self.assertEqual(
                detail.generated_by_username,
                self.operator.username,
            )
            self.assertEqual(
                detail.confirmed_by_username,
                self.super_admin.username,
            )

        self.db.expire_all()
        source_item = self.db.get(OrderItem, order_items[0].id)
        source_item.unit_cost_price = Decimal("13.00")
        source_item.line_cost_amount = Decimal("26.00")
        self.db.commit()
        with self.Session() as db:
            with self.assertRaisesRegex(RuntimeError, "来源订单项不一致"):
                get_supplier_settlement_detail(
                    db,
                    settlement_public_id=generated.settlement_public_id,
                )

    def test_builds_confirmed_workbook_from_safe_snapshots(self):
        self.supplier.name = "=测试供应商"
        self.product.name = "@测试商品"
        self.db.commit()
        self.add_order(
            suffix="REPORT-WORKBOOK",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 2),),
        )
        self.db.commit()
        generated = self.generate()
        self.confirm(generated.settlement_public_id)

        with self.Session() as db:
            detail = get_supplier_settlement_detail(
                db,
                settlement_public_id=generated.settlement_public_id,
            )
        workbook_stream = build_supplier_settlement_workbook(
            detail,
            exported_at=NOW + timedelta(minutes=2),
        )
        workbook = load_workbook(BytesIO(workbook_stream.getvalue()))

        self.assertEqual(workbook.sheetnames, ["结算汇总", "结算明细"])
        summary = workbook["结算汇总"]
        items = workbook["结算明细"]
        self.assertEqual(summary["B2"].value, generated.settlement_public_id)
        self.assertEqual(summary["B4"].value, "'=测试供应商")
        self.assertEqual(summary["B11"].value, 25)
        self.assertEqual(items.max_row, 2)
        self.assertEqual(items["B2"].value, "ORD-SETTLEMENT-REPORT-WORKBOOK")
        self.assertEqual(items["E2"].value, "'@测试商品")
        self.assertEqual(items["I2"].value, 12.5)
        self.assertEqual(items["J2"].value, 2)
        self.assertEqual(items["K2"].value, 25)

    def test_export_requires_confirmation_and_active_operations_role(self):
        self.add_order(
            suffix="REPORT-EXPORT",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()

        with self.assertRaisesRegex(
            ValueError,
            SETTLEMENT_EXPORT_STATE_MESSAGE,
        ):
            execute_supplier_settlement_export(
                self.engine,
                actor_admin_id=self.operator.id,
                settlement_public_id=generated.settlement_public_id,
                now=NOW + timedelta(minutes=2),
            )
        self.confirm(generated.settlement_public_id)

        for actor_id in (
            self.reviewer.id,
            self.inactive_operator.id,
            0,
        ):
            with self.subTest(actor_id=actor_id):
                with self.assertRaisesRegex(
                    PermissionError,
                    SETTLEMENT_EXPORT_PERMISSION_MESSAGE,
                ):
                    execute_supplier_settlement_export(
                        self.engine,
                        actor_admin_id=actor_id,
                        settlement_public_id=generated.settlement_public_id,
                        now=NOW + timedelta(minutes=2),
                    )

        exported = execute_supplier_settlement_export(
            self.engine,
            actor_admin_id=self.operator.id,
            settlement_public_id=generated.settlement_public_id,
            now=NOW + timedelta(minutes=2),
        )
        self.assertEqual(exported.status, "CONFIRMED")
        self.assertEqual(exported.total_cost_amount, Decimal("12.50"))
        self.assertTrue(exported.filename.endswith(".xlsx"))
        self.assertTrue(exported.content.startswith(b"PK"))
        with self.Session() as db:
            logs = db.query(AdminActionLog).filter_by(
                action_type="mall_supplier_settlement_export",
                target_id=generated.batch_id,
            ).all()
            self.assertEqual(len(logs), 1)
            self.assertEqual(logs[0].admin_id, self.operator.id)
            self.assertIn(generated.settlement_public_id, logs[0].description)
            self.assertIn("成本 12.50", logs[0].description)

    def test_export_fails_closed_for_tampered_confirmation_audit(self):
        self.add_order(
            suffix="REPORT-TAMPER",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()
        self.confirm(generated.settlement_public_id)
        self.db.expire_all()
        confirmation_log = self.db.query(AdminActionLog).filter_by(
            action_type="mall_supplier_settlement_confirm",
            target_id=generated.batch_id,
        ).one()
        confirmation_log.description = "篡改确认审计"
        self.db.commit()

        with self.assertRaisesRegex(RuntimeError, "确认审计证据不完整"):
            execute_supplier_settlement_export(
                self.engine,
                actor_admin_id=self.operator.id,
                settlement_public_id=generated.settlement_public_id,
                now=NOW + timedelta(minutes=2),
            )
        with self.Session() as db:
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_export"
                ).count(),
                0,
            )

    def test_export_rolls_back_audit_on_late_failure(self):
        self.add_order(
            suffix="REPORT-ROLLBACK",
            completed_at=PERIOD_START + timedelta(days=1),
            lines=((self.sku_one, 1),),
        )
        self.db.commit()
        generated = self.generate()
        self.confirm(generated.settlement_public_id)
        original_recorder = reporting_service.record_supplier_settlement_export

        def fail_after_audit(db, **request):
            original_recorder(db, **request)
            raise RuntimeError("模拟结算导出后失败")

        with patch.object(
            reporting_service,
            "record_supplier_settlement_export",
            side_effect=fail_after_audit,
        ):
            with self.assertRaisesRegex(RuntimeError, "模拟结算导出后失败"):
                execute_supplier_settlement_export(
                    self.engine,
                    actor_admin_id=self.operator.id,
                    settlement_public_id=generated.settlement_public_id,
                    now=NOW + timedelta(minutes=2),
                )
        with self.Session() as db:
            self.assertEqual(
                db.query(AdminActionLog).filter_by(
                    action_type="mall_supplier_settlement_export"
                ).count(),
                0,
            )


if __name__ == "__main__":
    unittest.main()
